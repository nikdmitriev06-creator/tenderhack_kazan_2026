import asyncio
import io
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.api.v1 import get_deps
from app.batch import (
    MAX_ROWS,
    BatchFormatError,
    BatchRow,
    BatchStore,
    BatchTooLarge,
    parse_upload,
    process_row,
    result_workbook,
    run_batch,
)
from app.catalog import Catalog
from app.extract import LLMItems
from app.glossary import Glossary
from app.main import create_app
from app.pipeline import Deps
from app.search import build_bm25
from app.session import SessionStore
from tests.fakes import FakeLLM, FakeRedis

FAKE = Path(__file__).parent / "fixtures" / "fake"
CATALOG = Catalog.load([FAKE / "ktru.json"])


def _deps(
    redis: FakeRedis | None = None, llm: FakeLLM | None = None, sent: list | None = None
) -> Deps:
    redis = redis or FakeRedis()

    async def enqueue(job_id: str) -> None:
        if sent is not None:
            sent.append(job_id)

    return Deps(
        catalog=CATALOG,
        bm25=build_bm25(CATALOG),
        vectors=None,
        llm=llm or FakeLLM([LLMItems(items=[])] * 50),  # type: ignore[arg-type]
        glossary=Glossary([]),
        sessions=SessionStore(redis, ttl_s=900),  # type: ignore[arg-type]
        min_score=0.2,
        batches=BatchStore(redis, ttl_s=900),  # type: ignore[arg-type]
        enqueue_batch=enqueue,
    )


def _names(rows: list[BatchRow]) -> list[str]:
    return [r.name for r in rows if not r.empty]


def test_upload_xlsx() -> None:
    rows = parse_upload("ste.xlsx", (FAKE / "ste.xlsx").read_bytes())
    assert len(_names(rows)) > 0


def test_upload_csv_in_cp1251_without_header_gives_the_same_names() -> None:
    xlsx = _names(parse_upload("ste.xlsx", (FAKE / "ste.xlsx").read_bytes()))
    cp1251 = _names(parse_upload("ste_cp1251.csv", (FAKE / "ste_cp1251.csv").read_bytes()))
    headerless = _names(
        parse_upload("ste_headerless.csv", (FAKE / "ste_headerless.csv").read_bytes())
    )
    assert xlsx == cp1251 == headerless


def test_upload_keeps_supplier_characteristics() -> None:
    rows = parse_upload("ste.xlsx", (FAKE / "ste.xlsx").read_bytes())
    assert any(r.characteristics for r in rows)


def test_blank_row_is_kept_as_skipped() -> None:
    # Строка без наименования, но с описанием, не теряется, а помечается пропущенной.
    # Полностью пустые строки читатель таблиц отбрасывает сам: это не данные.
    data = (
        "Наименование;Описание\nБумага офисная А4;\n;без наименования\nПерчатки нитриловые M;\n"
    ).encode()
    rows = parse_upload("cards.csv", data)
    assert [(r.row, r.empty) for r in rows] == [(1, False), (2, True), (3, False)]


def test_file_without_name_column_is_rejected() -> None:
    with pytest.raises(BatchFormatError):
        parse_upload("cards.csv", "Цена;Количество\n100;2\n".encode())


def test_too_many_rows_is_rejected() -> None:
    data = ("Наименование\n" + "Бумага офисная А4\n" * (MAX_ROWS + 1)).encode()
    with pytest.raises(BatchTooLarge, match=str(MAX_ROWS)):
        parse_upload("cards.csv", data)


def test_row_failure_does_not_stop_the_batch() -> None:
    class BrokenBM25:
        def search(self, *_: object, **__: object) -> list:
            raise RuntimeError("индекс сломан")

    deps = _deps()
    deps.bm25 = BrokenBM25()  # type: ignore[assignment]
    result = asyncio.run(process_row(BatchRow(row=1, name="Бумага"), deps))
    assert result.status == "failed"
    assert result.error is not None and "индекс сломан" in result.error


def test_empty_row_is_skipped() -> None:
    result = asyncio.run(process_row(BatchRow(row=2, name="", empty=True), _deps()))
    assert result.status == "skipped"


def test_run_batch_processes_every_row_and_finishes() -> None:
    redis = FakeRedis()
    deps = _deps(redis)
    rows = [
        BatchRow(row=1, name="Батарейка AAA щелочная 1,5 В"),
        BatchRow(row=2, name="", empty=True),
        BatchRow(row=3, name="Услуги по уборке помещений"),
    ]
    store = deps.batches
    assert store is not None
    job_id = asyncio.run(store.create("cards.xlsx", rows))
    asyncio.run(run_batch(job_id, deps))
    state = asyncio.run(store.state(job_id))
    assert state is not None
    assert (state.status, state.total, state.processed) == ("done", 3, 3)
    assert [r.status for r in state.rows][1] == "skipped"
    assert state.rows[2].ktru_code is None


def test_result_workbook_adds_module_columns() -> None:
    rows = [BatchRow(row=1, name="Бумага", raw={"Наименование": "Бумага", "Цена": "100"})]
    from app.batch import BatchRowResult, BatchState

    state = BatchState(
        job_id="x",
        status="done",
        filename="cards.xlsx",
        total=1,
        processed=1,
        rows=[
            BatchRowResult(
                row=1, name="Бумага", status="needs_answer", ktru_code=None, filled=2, issues=[]
            )
        ],
    )
    wb = load_workbook(io.BytesIO(result_workbook(state, rows)))
    ws = wb.active
    header = [c.value for c in ws[1]]
    assert header == ["Наименование", "Цена", "Код КТРУ", "Заполнено", "Замечания", "Статус"]
    assert ws.max_row == 2


def test_api_accepts_file_and_queues_job() -> None:
    sent: list[str] = []
    redis = FakeRedis()
    app = create_app()
    app.dependency_overrides[get_deps] = lambda: _deps(redis, sent=sent)
    client = TestClient(app)
    data = "Наименование\nБумага офисная А4\n".encode()
    r = client.post("/api/v1/batch", files={"file": ("cards.csv", data, "text/csv")})
    assert r.status_code == 202
    job_id = r.json()["job_id"]
    assert sent == [job_id]
    state = client.get(f"/api/v1/batch/{job_id}").json()
    assert (state["status"], state["total"]) == ("queued", 1)
    assert client.get(f"/api/v1/batch/{job_id}/result.xlsx").status_code == 409


def test_api_rejects_oversized_file_with_413() -> None:
    app = create_app()
    app.dependency_overrides[get_deps] = lambda: _deps()
    client = TestClient(app)
    data = ("Наименование\n" + "Бумага офисная А4\n" * (MAX_ROWS + 1)).encode()
    r = client.post("/api/v1/batch", files={"file": ("cards.csv", data, "text/csv")})
    assert r.status_code == 413
    assert str(MAX_ROWS) in r.json()["detail"]


def test_api_unknown_batch_is_404() -> None:
    app = create_app()
    app.dependency_overrides[get_deps] = lambda: _deps()
    assert TestClient(app).get("/api/v1/batch/nope").status_code == 404


def test_row_out_of_catalog_is_not_found_with_the_reason() -> None:
    # Ревью: строка вне справочника была «нужен ответ», хотя спрашивать нечего.
    result = asyncio.run(process_row(BatchRow(row=3, name="Услуги по уборке помещений"), _deps()))
    assert result.status == "not_found"
    assert result.note is not None and result.note.startswith("Категория не найдена")


def test_result_workbook_explains_not_found_rows() -> None:
    from app.batch import BatchRowResult, BatchState

    rows = [BatchRow(row=1, name="Уборка", raw={"Наименование": "Уборка"})]
    note = "Категория не найдена: в справочнике нет подходящей позиции."
    state = BatchState(
        job_id="x",
        status="done",
        filename="cards.xlsx",
        total=1,
        processed=1,
        rows=[
            BatchRowResult(
                row=1, name="Уборка", status="not_found", ktru_code=None, filled=0, issues=[],
                note=note,
            )
        ],
    )  # fmt: skip
    ws = load_workbook(io.BytesIO(result_workbook(state, rows))).active
    assert [c.value for c in ws[2]][-2:] == [note, "нет в справочнике"]
